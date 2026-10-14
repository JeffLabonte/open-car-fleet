import tempfile
import uuid
from io import BytesIO, StringIO
from pathlib import Path

from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from shop.exporters import export_fleet_to_excel
from shop.models.attachment import Attachment
from shop.models.car import Car
from shop.models.fleet import Fleet, FleetMembership
from shop.models.job import WorkJob
from shop.models.report import Report
from shop.models.user import ShopUser


class FleetExportServiceTests(TestCase):
    def setUp(self) -> None:
        self.owner = ShopUser.objects.create_user(
            username="export-owner",
            email="export-owner@example.com",
            password="pass1234",
            is_mechanic=True,
        )
        self.fleet = Fleet.objects.create(name="Primary Fleet", created_by=self.owner)
        FleetMembership.objects.create(
            fleet=self.fleet,
            user=self.owner,
            role=FleetMembership.ROLE_OWNER,
        )

        self.other_fleet = Fleet.objects.create(name="Other Fleet", created_by=self.owner)
        FleetMembership.objects.create(
            fleet=self.other_fleet,
            user=self.owner,
            role=FleetMembership.ROLE_OWNER,
        )

        self.primary_car = Car.objects.create(
            fleet=self.fleet,
            usual_name="Daily Driver",
            make="Toyota",
            model="Corolla",
            vin="2T1BURHE5JC512345",
        )
        self.other_car = Car.objects.create(
            fleet=self.other_fleet,
            usual_name="Spare Car",
            make="Honda",
            model="Civic",
            vin="2HGFG12698H512345",
        )

        WorkJob.objects.create(
            car=self.primary_car,
            title="Oil Change",
            assigned_to=self.owner,
            required_items=["Oil", "Filter"],
            status="pending",
            urgency="soon",
        )
        WorkJob.objects.create(
            car=self.other_car,
            title="Do Not Export",
        )

        brake_report = Report.objects.create(
            car=self.primary_car,
            mileage=120000,
            job_name="Brake Service",
            assigned_to=self.owner,
            date_done=timezone.now().date(),
        )
        ContentType.objects.get_for_model(brake_report)
        Attachment.objects.create(
            content_type=ContentType.objects.get_for_model(brake_report),
            object_id=str(brake_report.pk),
            source_type=Attachment.SOURCE_UPLOAD,
            file=SimpleUploadedFile("invoice.pdf", b"%PDF-1.4 invoice", content_type="application/pdf"),
            display_name="invoice.pdf",
            mime_type="application/pdf",
            kind="document",
        )
        Attachment.objects.create(
            content_type=ContentType.objects.get_for_model(brake_report),
            object_id=str(brake_report.pk),
            source_type=Attachment.SOURCE_EXTERNAL,
            url="https://example.com/before.jpg",
            display_name="before.jpg",
            kind="link",
        )
        Report.objects.create(
            car=self.other_car,
            job_name="Skip Report",
            date_done=timezone.now().date(),
        )

    def test_export_fleet_to_excel_includes_expected_sheets_and_rows(self):
        workbook_file = export_fleet_to_excel(self.fleet)
        workbook = load_workbook(filename=BytesIO(workbook_file.content))

        self.assertIn("meta", workbook.sheetnames)
        self.assertIn("fleet", workbook.sheetnames)
        self.assertIn("memberships", workbook.sheetnames)
        self.assertIn("cars_import", workbook.sheetnames)
        self.assertIn("workjobs_import", workbook.sheetnames)
        self.assertIn("reports_import", workbook.sheetnames)

        cars_rows = list(workbook["cars_import"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(cars_rows), 1)
        self.assertEqual(cars_rows[0][3], "Toyota")
        self.assertNotIn("2HGFG12698H512345", [row[7] for row in cars_rows])

        job_rows = list(workbook["workjobs_import"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(job_rows), 1)
        self.assertEqual(job_rows[0][3], "Oil Change")
        self.assertEqual(job_rows[0][14], "Oil\nFilter")

        report_rows = list(workbook["reports_import"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(report_rows), 1)
        self.assertEqual(report_rows[0][4], "Brake Service")
        self.assertEqual(report_rows[0][12], "invoice.pdf\nhttps://example.com/before.jpg")


class FleetExportViewTests(TestCase):
    def setUp(self) -> None:
        self.owner = ShopUser.objects.create_user(
            username="export-view-owner",
            email="export-view-owner@example.com",
            password="pass1234",
            is_mechanic=True,
        )
        self.manager = ShopUser.objects.create_user(
            username="export-view-manager",
            email="export-view-manager@example.com",
            password="pass1234",
        )
        self.member = ShopUser.objects.create_user(
            username="export-view-member",
            email="export-view-member@example.com",
            password="pass1234",
        )

        self.fleet = Fleet.objects.create(name="Export View Fleet", created_by=self.owner)
        FleetMembership.objects.create(
            fleet=self.fleet,
            user=self.owner,
            role=FleetMembership.ROLE_OWNER,
        )
        FleetMembership.objects.create(
            fleet=self.fleet,
            user=self.manager,
            role=FleetMembership.ROLE_MANAGER,
        )
        FleetMembership.objects.create(
            fleet=self.fleet,
            user=self.member,
            role=FleetMembership.ROLE_MEMBER,
        )

        self.primary_car = Car.objects.create(
            fleet=self.fleet,
            usual_name="Exportable",
            make="Ford",
            model="Focus",
            vin="1FAHP3F28CL512345",
        )

        self.other_fleet = Fleet.objects.create(name="External Fleet", created_by=self.owner)
        FleetMembership.objects.create(
            fleet=self.other_fleet,
            user=self.owner,
            role=FleetMembership.ROLE_OWNER,
        )
        Car.objects.create(
            fleet=self.other_fleet,
            usual_name="Hidden",
            make="Tesla",
            model="Model 3",
            vin="5YJ3E1EA7LF512345",
        )

    def test_unauthenticated_user_is_redirected_from_fleet_export(self):
        response = self.client.get(reverse("shop-fleet-export", args=[self.fleet.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].startswith(f"{reverse('shop-login')}?next="))

    def test_non_manager_is_redirected_from_fleet_export(self):
        self.client.force_login(self.member)

        response = self.client.get(reverse("shop-fleet-export", args=[self.fleet.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("shop-fleet-detail", args=[self.fleet.pk]))

    def test_manager_can_download_fleet_export_workbook(self):
        self.client.force_login(self.manager)

        response = self.client.get(reverse("shop-fleet-export", args=[self.fleet.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('attachment; filename="fleet_export-view-fleet_', response["Content-Disposition"])

        workbook = load_workbook(filename=BytesIO(response.content))
        cars_rows = list(workbook["cars_import"].iter_rows(min_row=2, values_only=True))
        vins = [row[7] for row in cars_rows]
        self.assertIn("1FAHP3F28CL512345", vins)
        self.assertNotIn("5YJ3E1EA7LF512345", vins)


class ExportFleetCommandTests(TestCase):
    def setUp(self) -> None:
        self.user = ShopUser.objects.create_user(
            username="command-owner",
            email="command-owner@example.com",
            password="pass1234",
        )
        self.fleet = Fleet.objects.create(name="Command Fleet", created_by=self.user)
        FleetMembership.objects.create(
            fleet=self.fleet,
            user=self.user,
            role=FleetMembership.ROLE_OWNER,
        )

    def test_export_fleet_command_writes_excel_file(self):
        Car.objects.create(
            fleet=self.fleet,
            usual_name="Command Export Car",
            make="Mazda",
            model="3",
            vin="JM1BK323171512345",
        )
        output = StringIO()

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            output_path = Path(handle.name)
        output_path.unlink(missing_ok=True)

        try:
            call_command(
                "export_fleet",
                str(self.fleet.pk),
                "--output",
                str(output_path),
                stdout=output,
            )
            self.assertTrue(output_path.exists())
            workbook = load_workbook(filename=str(output_path))
            self.assertIn("cars_import", workbook.sheetnames)
            cars_sheet = workbook["cars_import"]
            rows = list(cars_sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0][3], "Mazda")
        finally:
            output_path.unlink(missing_ok=True)

    def test_export_fleet_command_errors_for_unknown_fleet(self):
        with self.assertRaises(CommandError):
            call_command("export_fleet", str(uuid.uuid4()))

    def test_export_fleet_command_rejects_invalid_fleet_uuid(self):
        with self.assertRaises(CommandError):
            call_command("export_fleet", "not-a-uuid")
