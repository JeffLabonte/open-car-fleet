import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.test.utils import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from shop.exporters import export_garage_to_excel
from shop.forms import ReportForm
from shop.models.car import Car
from shop.models.garage import Garage
from shop.models.garage import GarageMembership
from shop.models.user import ShopUser
from shop.tests.helpers import PNG_SIGNATURE


class SecurityHardeningTests(TestCase):
    """Regression tests for security audit fixes."""

    def setUp(self) -> None:
        self.owner = ShopUser.objects.create_user(
            username='security-owner',
            email='security-owner@example.com',
            password='pass1234',
        )
        self.garage = Garage.objects.create(name='Security Garage', created_by=self.owner)
        GarageMembership.objects.create(
            garage=self.garage,
            user=self.owner,
            role=GarageMembership.ROLE_OWNER,
        )
        self.car = Car.objects.create(
            garage=self.garage,
            make='Toyota',
            model='Yaris',
            vin='JTDKB20U793512349',
        )

    def test_login_page_rejects_javascript_next_url(self):
        response = self.client.get(reverse('shop-login'), {'next': 'javascript:alert(document.cookie)'})
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'javascript:alert')
        self.assertContains(response, '"next-url"')

    def test_theme_view_rejects_absolute_next_url(self):
        response = self.client.get(reverse('shop-theme', args=['light']), {'next': 'https://evil.example.com'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('shop-index'))

    def test_theme_view_allows_local_next_url(self):
        response = self.client.get(reverse('shop-theme', args=['dark']), {'next': '/cars/'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, '/cars/')

    def test_report_attachment_rejects_disallowed_extension(self):
        form = ReportForm(
            data={
                'mileage': '',
                'job_name': 'Oil change',
                'date_done': '2026-01-15',
                'note': '',
                'additional_information': '',
            },
            files={
                'attachments': SimpleUploadedFile(
                    'payload.html',
                    b'<html></html>',
                    content_type='image/jpeg',
                ),
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('attachments', form.errors)

    def test_report_attachment_rejects_spoofed_magic_bytes(self):
        form = ReportForm(
            data={
                'mileage': '',
                'job_name': 'Oil change',
                'date_done': '2026-01-15',
                'note': '',
                'additional_information': '',
            },
            files={
                'attachments': SimpleUploadedFile(
                    'payload.png',
                    b'<html><script>alert(1)</script></html>',
                    content_type='image/png',
                ),
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('attachments', form.errors)

    def test_report_attachment_rejects_oversized_upload(self):
        from unittest.mock import patch

        from shop.forms.base import AttachmentField

        # Patching the shared limit keeps this test cheap regardless of the
        # production ceiling (500 MB).
        with patch.object(AttachmentField, 'max_upload_bytes', 10):
            form = ReportForm(
                data={
                    'mileage': '',
                    'job_name': 'Oil change',
                    'date_done': '2026-01-15',
                    'note': '',
                    'additional_information': '',
                },
                files={
                    'attachments': SimpleUploadedFile(
                        'large.png',
                        PNG_SIGNATURE + b'0' * 11,
                        content_type='image/png',
                    ),
                },
                user=self.owner,
                garage=self.garage,
            )
            self.assertFalse(form.is_valid())
        self.assertIn('attachments', form.errors)

    def test_report_attachment_rejects_too_many_files(self):
        from shop.forms.base import ATTACHMENT_MAX_FILES

        form = ReportForm(
            data={
                'mileage': '',
                'job_name': 'Oil change',
                'date_done': '2026-01-15',
                'note': '',
                'additional_information': '',
            },
            files={
                'attachments': [
                    SimpleUploadedFile(f'photo-{index}.png', PNG_SIGNATURE, content_type='image/png')
                    for index in range(ATTACHMENT_MAX_FILES + 1)
                ],
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('attachments', form.errors)

    def test_set_test_session_is_disabled_when_debug_is_false(self):
        response = self.client.get(reverse('shop-set-test-session'), {'email': 'e2e@example.com'})
        self.assertEqual(response.status_code, 404)

    @override_settings(ROOT_URLCONF='settings.urls')
    def test_set_test_session_is_not_present_in_production_urls(self):
        response = self.client.get('/set-test-session/', {'email': 'e2e@example.com'})
        self.assertEqual(response.status_code, 404)

    @override_settings(DEBUG=True)
    def test_set_test_session_signs_user_in_when_debug(self):
        response = self.client.get(
            reverse('shop-set-test-session'),
            {'email': 'e2e-user@example.com', 'next': '/cars/'},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, '/cars/')
        self.assertTrue(self.client.session.get('_auth_user_id'))

    @override_settings(DEBUG=True)
    def test_set_test_session_requires_email(self):
        response = self.client.get(reverse('shop-set-test-session'), {'next': '/cars/'})
        self.assertEqual(response.status_code, 400)

    def test_excel_export_neutralizes_formula_injection(self):
        self.garage.name = "=cmd|'/c calc'!A0"
        self.garage.save(update_fields=['name'])
        workbook_file = export_garage_to_excel(self.garage)
        workbook = load_workbook(filename=BytesIO(workbook_file.content))
        garage_name_cell = workbook['garage'].cell(row=2, column=2).value
        self.assertEqual(garage_name_cell, "'=cmd|'/c calc'!A0")

    def test_excel_export_neutralizes_whitespace_prefixed_formulas(self):
        self.garage.name = "\t\n=cmd|'/c calc'!A0"
        self.garage.save(update_fields=['name'])
        workbook_file = export_garage_to_excel(self.garage)
        workbook = load_workbook(filename=BytesIO(workbook_file.content))
        garage_name_cell = workbook['garage'].cell(row=2, column=2).value
        self.assertEqual(garage_name_cell, "'\t\n=cmd|'/c calc'!A0")

    def test_known_shop_create_invalid_post_returns_form_not_500(self):
        self.client.force_login(self.owner)
        response = self.client.post(reverse('shop-known-shop-create'), data={'name': ''})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Add shop')

    def test_report_form_rejects_too_many_external_links(self):
        from shop.forms.base import MAX_EXTERNAL_LINKS

        links = '\n'.join(f'https://example.com/{index}' for index in range(MAX_EXTERNAL_LINKS + 1))
        form = ReportForm(
            data={
                'mileage': '100000',
                'job_name': 'Brake service',
                'date_done': '2026-01-15',
                'external_links': links,
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('external_links', form.errors)

    def test_report_form_rejects_external_link_too_long(self):
        from shop.forms.base import MAX_EXTERNAL_LINK_LENGTH

        long_link = 'https://example.com/' + 'a' * MAX_EXTERNAL_LINK_LENGTH
        form = ReportForm(
            data={
                'mileage': '100000',
                'job_name': 'Brake service',
                'date_done': '2026-01-15',
                'external_links': long_link,
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('external_links', form.errors)

    def test_report_form_rejects_too_many_staged_attachment_ids(self):
        from shop.forms.base import STAGED_ATTACHMENT_MAX_IDS

        ids = ','.join(str(index) for index in range(STAGED_ATTACHMENT_MAX_IDS + 1))
        form = ReportForm(
            data={
                'mileage': '100000',
                'job_name': 'Brake service',
                'date_done': '2026-01-15',
                'staged_attachments': ids,
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('staged_attachments', form.errors)

    def test_report_form_rejects_malformed_staged_attachment_ids(self):
        form = ReportForm(
            data={
                'mileage': '100000',
                'job_name': 'Brake service',
                'date_done': '2026-01-15',
                'staged_attachments': '1,abc,99999999999999999999',
            },
            user=self.owner,
            garage=self.garage,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('staged_attachments', form.errors)

    def test_production_security_defaults_are_enforced(self):
        project_root = Path(__file__).resolve().parents[3]
        src_dir = project_root / 'src'
        env = {
            **os.environ,
            'DJANGO_SECRET_KEY': 'production-test-secret-key-not-used-in-deployment',
            'DEBUG': 'False',
            'HANKO_API_URL': 'https://hanko.example.com',
            'DJANGO_ALLOWED_HOSTS': 'fleet.example.com',
        }
        # Clear any local .env so the test uses the env vars above.
        env.pop('POSTGRES_DB', None)

        code = (
            "import os; "
            "os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'settings.settings'); "
            "import django; django.setup(); "
            "from django.conf import settings; "
            "print(f'ssl_redirect={settings.SECURE_SSL_REDIRECT}'); "
            "print(f'hsts_seconds={settings.SECURE_HSTS_SECONDS}'); "
            "print(f'session_cookie_secure={settings.SESSION_COOKIE_SECURE}'); "
            "print(f'csrf_cookie_secure={settings.CSRF_COOKIE_SECURE}');"
        )
        result = subprocess.run(
            [sys.executable, '-c', code],
            cwd=src_dir,
            env=env,
            capture_output=True,
            text=True,
            check=True,
        )
        output = result.stdout.strip().splitlines()
        values = dict(line.split('=', 1) for line in output)

        self.assertEqual(values['ssl_redirect'], 'True')
        self.assertEqual(values['hsts_seconds'], '31536000')
        self.assertEqual(values['session_cookie_secure'], 'True')
        self.assertEqual(values['csrf_cookie_secure'], 'True')

    def test_development_security_defaults_are_relaxed(self):
        project_root = Path(__file__).resolve().parents[3]
        src_dir = project_root / 'src'
        env = {
            **os.environ,
            'DJANGO_SECRET_KEY': 'development-test-secret-key-not-used-in-deployment',
            'DEBUG': 'True',
            'HANKO_API_URL': 'https://hanko.example.com',
        }
        env.pop('POSTGRES_DB', None)

        code = (
            "import os; "
            "os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'settings.settings'); "
            "import django; django.setup(); "
            "from django.conf import settings; "
            "print(f'ssl_redirect={settings.SECURE_SSL_REDIRECT}'); "
            "print(f'hsts_seconds={settings.SECURE_HSTS_SECONDS}'); "
            "print(f'session_cookie_secure={settings.SESSION_COOKIE_SECURE}'); "
            "print(f'csrf_cookie_secure={settings.CSRF_COOKIE_SECURE}');"
        )
        result = subprocess.run(
            [sys.executable, '-c', code],
            cwd=src_dir,
            env=env,
            capture_output=True,
            text=True,
            check=True,
        )
        output = result.stdout.strip().splitlines()
        values = dict(line.split('=', 1) for line in output)

        self.assertEqual(values['ssl_redirect'], 'False')
        self.assertEqual(values['hsts_seconds'], '0')
        self.assertEqual(values['session_cookie_secure'], 'False')
        self.assertEqual(values['csrf_cookie_secure'], 'False')
